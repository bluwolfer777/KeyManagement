"""Ricevuta di prestito in PDF, disegnata a partire dal modello Excel.

Dal file modelli/ricevuta_chiavi.xlsx vengono letti testi, caratteri, allineamenti, bordi,
riempimenti, larghezze di colonna, altezze di riga, margini, scala di stampa, logo e linea
decorativa; il PDF li riproduce e riempie i campi:
  - cella a destra di "Nome dell'autore"  -> operatore
  - cella a destra di "Cognome e Nome"    -> utente
  - righe della tabella sotto le intestazioni "Data", "Tipo di chiave", "Firma di ricevuta",
    "Firma di ritorno", "Data di ritorno" -> una riga per chiave.
Se le righe non bastano si aggiungono pagine con la stessa intestazione.
"""
import base64
import io
import os
import re
import zipfile
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter
from PIL import Image
from reportlab.lib.colors import Color, black
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfgen import canvas

import config

TEMPLATE = os.environ.get("GC_RICEVUTA_TEMPLATE", os.path.join(config.BASE_DIR, "modelli", "ricevuta_chiavi.xlsx"))
EMU_PT = 12700
BORDER_WIDTH = {"hair": .25, "thin": .5, "dotted": .5, "dashed": .5, "medium": 1.5,
                "mediumDashed": 1.5, "double": 1.5, "thick": 2.25}
_cache = {}


# --- Lettura del modello -------------------------------------------------------

def _col_pt(chars):
    return int(chars * 7 + 5) * .75          # larghezza Excel (caratteri) -> punti


def _rgb(color, default=black):
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and re.fullmatch(r"[0-9A-Fa-f]{8}", rgb) and getattr(color, "type", "rgb") == "rgb":
        return Color(int(rgb[2:4], 16) / 255, int(rgb[4:6], 16) / 255, int(rgb[6:8], 16) / 255)
    return default


def _anchor(xml, tag):
    m = re.search(rf"<xdr:{tag}>(.*?)</xdr:{tag}>", xml, re.S)
    return {k: int(v) for k, v in re.findall(r"<xdr:(\w+)>(-?\d+)</xdr:\w+>", m.group(1))}


class Model:
    def __init__(self, path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        if ws.print_area:
            area = ws.print_area[0] if isinstance(ws.print_area, (list, tuple)) else ws.print_area
            area = area.split("!")[-1].replace("$", "")
        else:
            area = ws.dimensions
        first, last = area.split(":")
        c1, r1 = openpyxl.utils.cell.coordinate_from_string(first)
        c2, r2 = openpyxl.utils.cell.coordinate_from_string(last)
        self.cols = range(openpyxl.utils.column_index_from_string(c1), openpyxl.utils.column_index_from_string(c2) + 1)
        self.rows = range(r1, r2 + 1)

        default_w = ws.sheet_format.defaultColWidth or 8.43
        widths = {c: default_w for c in self.cols}
        for dim in ws.column_dimensions.values():
            if dim.width and dim.min:
                for c in range(dim.min, (dim.max or dim.min) + 1):
                    widths[c] = dim.width
        self.col_x = {}
        x = 0
        for c in list(self.cols) + [self.cols[-1] + 1]:
            self.col_x[c] = x
            x += _col_pt(widths.get(c, default_w))
        default_h = ws.sheet_format.defaultRowHeight or 15
        self.row_y = {}
        y = 0
        for r in list(self.rows) + [self.rows[-1] + 1]:
            self.row_y[r] = y
            y += ws.row_dimensions[r].height or default_h
        self.width = self.col_x[self.cols[-1] + 1]
        self.height = self.row_y[self.rows[-1] + 1]

        self.merged = {}
        self.hidden = set()
        self.span = {}          # cella -> (riga1, col1, riga2, col2) dell'eventuale unione
        for rng in ws.merged_cells.ranges:
            self.merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.span[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
                    if (r, c) != (rng.min_row, rng.min_col):
                        self.hidden.add((r, c))
        self.cells = {(r, c): ws.cell(r, c) for r in self.rows for c in self.cols}

        ps = ws.page_setup
        self.scale = (ps.scale or 100) / 100
        self.pagesize = A3 if str(ps.paperSize) == "8" else A4
        if ps.orientation == "landscape":
            self.pagesize = landscape(self.pagesize)
        m = ws.page_margins
        self.margin_left, self.margin_top = (m.left or .7) * 72, (m.top or .75) * 72
        self.margin_bottom = (m.bottom or .75) * 72

        self._find_fields()
        self._read_drawings(path)

    def rect(self, r, c):
        r2, c2 = self.merged.get((r, c), (r, c))
        return self.col_x[c], self.row_y[r], self.col_x[c2 + 1] - self.col_x[c], self.row_y[r2 + 1] - self.row_y[r]

    def _text(self, r, c):
        v = self.cells.get((r, c))
        return str(v.value).strip().lower() if v is not None and v.value is not None else ""

    def _right_of(self, r, c):
        _, c2 = self.merged.get((r, c), (r, c))
        return (r, c2 + 1) if c2 + 1 in self.cols else None

    def _find_fields(self):
        self.fields = {}
        self.columns = {}
        self.header_row = None
        for (r, c), _ in sorted(self.cells.items()):
            t = self._text(r, c)
            if not t:
                continue
            if t.startswith("nome dell'autore") or t.startswith("autore"):
                self.fields["operator"] = self._right_of(r, c)
            elif t.startswith("cognome e nome") or t.startswith("nome e cognome"):
                self.fields["name"] = self._right_of(r, c)
            elif "firma" in t and "ricevut" in t:
                self.columns["sig_out"], self.header_row = c, r
            elif "firma" in t and "ritorn" in t:
                self.columns["sig_in"] = c
            elif "data" in t and "ritorn" in t:
                self.columns["date_in"] = c
            elif t == "data":
                self.columns["date_out"] = c
            elif "chiav" in t:
                self.columns["key"] = c
        if self.header_row is None:
            raise ValueError("Nel modello non trovo l'intestazione 'Firma di ricevuta'")
        anchor = self.columns.get("date_out", self.cols[0])
        start = self.header_row + 1
        self.body = []
        r = start
        while r in self.rows:
            r2, _ = self.merged.get((r, anchor), (r, anchor))
            self.body.append((r, r2))
            r = r2 + 1

    def _read_drawings(self, path):
        self.logo = None
        self.line_group = []
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            drawing = next((n for n in names if n.startswith("xl/drawings/drawing") and n.endswith(".xml")), None)
            if not drawing:
                return
            xml = z.read(drawing).decode("utf-8")
            rels_name = drawing.replace("drawings/", "drawings/_rels/") + ".rels"
            rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(rels_name).decode())) \
                if rels_name in names else {}
            for m in re.finditer(r"<xdr:twoCellAnchor[^>]*>(.*?)</xdr:twoCellAnchor>", xml, re.S):
                body = m.group(1)
                fr, to = _anchor(body, "from"), _anchor(body, "to")
                x1 = self.col_x.get(fr["col"] + 1, 0) + fr["colOff"] / EMU_PT
                y1 = self.row_y.get(fr["row"] + 1, 0) + fr["rowOff"] / EMU_PT
                x2 = self.col_x.get(to["col"] + 1, self.width) + to["colOff"] / EMU_PT
                y2 = self.row_y.get(to["row"] + 1, self.height) + to["rowOff"] / EMU_PT
                if "<xdr:pic>" in body:
                    rid = re.search(r'r:embed="([^"]+)"', body).group(1)
                    target = os.path.normpath(os.path.join("xl/drawings", rels[rid])).replace("\\", "/")
                    ext = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"/>', body)
                    w, h = (int(ext.group(1)) / EMU_PT, int(ext.group(2)) / EMU_PT) if ext else (x2 - x1, y2 - y1)
                    self.logo = (z.read(target), x1, y1, w, h)
                elif "<xdr:grpSp>" in body:
                    ch = re.search(r'<a:chOff x="(\d+)" y="(\d+)"/><a:chExt cx="(\d+)" cy="(\d+)"/>', body)
                    chx, chw = int(ch.group(1)), int(ch.group(3)) or 1
                    for sp in re.finditer(r'<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/></a:xfrm>'
                                          r'<a:prstGeom prst="line">.*?<a:ln w="(\d+)"', body, re.S):
                        sx, _, sw, _, lw = map(int, sp.groups())
                        fx = lambda v: x1 + (v - chx) / chw * (x2 - x1)
                        self.line_group.append((fx(sx), fx(sx + sw), (y1 + y2) / 2, int(lw) / EMU_PT))


def model():
    mtime = os.path.getmtime(TEMPLATE)
    if _cache.get("mtime") != mtime:
        _cache.update(mtime=mtime, model=Model(TEMPLATE))
    return _cache["model"]


# --- Disegno ---------------------------------------------------------------------

def _font(cell, size_default=10):
    f = cell.font
    bold = bool(f and f.b)
    return ("Helvetica-Bold" if bold else "Helvetica"), (f.sz if f and f.sz else size_default), _rgb(f.color if f else None)


def _draw_text(c, text, x, y, w, h, font, size, color, halign="center", valign="center", wrap=True):
    lines = simpleSplit(text, font, size, w - 4) if wrap else [text]
    lead = size * 1.15
    while len(lines) * lead > h and size > 6:
        size -= .5
        lead = size * 1.15
        lines = simpleSplit(text, font, size, w - 4) if wrap else [text]
    block = len(lines) * lead
    top = y + (h - block) / 2 if valign == "center" else (y + h - block - 2 if valign == "bottom" else y + 2)
    c.setFont(font, size)
    c.setFillColor(color)
    for i, line in enumerate(lines):
        base = -(top + lead * i + size * .85)
        if halign == "left":
            c.drawString(x + 2, base, line)
        elif halign == "right":
            c.drawRightString(x + w - 2, base, line)
        else:
            c.drawCentredString(x + w / 2, base, line)


def _signature_reader(data_url):
    if not data_url or "," not in data_url:
        return None
    try:
        img = Image.open(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1]))).convert("RGBA")
    except Exception:
        return None
    bbox = img.getchannel("A").getbbox()
    if bbox:
        pad = 6
        img = img.crop((max(bbox[0] - pad, 0), max(bbox[1] - pad, 0),
                        min(bbox[2] + pad, img.width), min(bbox[3] + pad, img.height)))
    flat = Image.new("RGB", img.size, "white")
    flat.paste(img, mask=img.getchannel("A"))
    return ImageReader(flat)


def _draw_image(c, reader, x, y, w, h, pad=3):
    iw, ih = reader.getSize()
    k = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    dw, dh = iw * k, ih * k
    c.drawImage(reader, x + (w - dw) / 2, -(y + (h + dh) / 2), dw, dh)


def _draw_page(c, m, header, rows):
    c.saveState()
    c.translate(m.margin_left, m.pagesize[1] - m.margin_top)
    c.scale(m.scale, m.scale)
    # Riempimenti e testi del modello
    for (r, col), cell in m.cells.items():
        if (r, col) in m.hidden:
            continue
        x, y, w, h = m.rect(r, col)
        fill = cell.fill
        if fill is not None and fill.fill_type == "solid":
            c.setFillColor(_rgb(fill.fgColor, None) or Color(.85, .85, .85))
            c.rect(x, -(y + h), w, h, stroke=0, fill=1)
        if cell.value is not None:
            font, size, color = _font(cell)
            al = cell.alignment
            _draw_text(c, str(cell.value).strip(), x, y, w, h, font, size, color,
                       (al.horizontal or "left") if al else "left",
                       "center" if al and al.vertical == "center" else "bottom",
                       wrap=bool(al and al.wrap_text) or r == m.header_row)
    # Bordi
    c.setStrokeColor(black)
    c.setLineCap(2)
    for (r, col), cell in m.cells.items():
        b = cell.border
        if b is None:
            continue
        x, y = m.col_x[col], m.row_y[r]
        x2, y2 = m.col_x[col + 1], m.row_y[r + 1]
        r1, c1, r2, c2 = m.span.get((r, col), (r, col, r, col))
        # Nelle celle unite Excel disegna solo il contorno esterno dell'unione.
        outer = {"top": r == r1, "bottom": r == r2, "left": col == c1, "right": col == c2}
        for side, (ax, ay, bx, by) in (("top", (x, y, x2, y)), ("bottom", (x, y2, x2, y2)),
                                       ("left", (x, y, x, y2)), ("right", (x2, y, x2, y2))):
            style = getattr(b, side).style
            if style and outer[side]:
                c.setLineWidth(BORDER_WIDTH.get(style, .5))
                c.line(ax, -ay, bx, -by)
    # Linea decorativa e, sopra, il logo (stesso ordine di Excel: il logo copre i primi quadratini)
    c.setLineCap(0)
    for x1, x2, y, lw in m.line_group:
        c.setLineWidth(lw)
        c.line(x1, -y, x2, -y)
    if m.logo:
        data, lx, ly, lw, lh = m.logo
        c.drawImage(ImageReader(io.BytesIO(data)), lx, -(ly + lh), lw, lh)
    # Campi dell'intestazione
    for key, value in header.items():
        pos = m.fields.get(key)
        if pos and value:
            x, y, w, h = m.rect(*pos)
            if key in ("operator", "name"):
                w = m.width - x
            _draw_text(c, value, x, y, w, h, "Helvetica", 10, black, "left", "center")
    # Righe della tabella
    for (r1, r2), row in zip(m.body, rows):
        y, h = m.row_y[r1], m.row_y[r2 + 1] - m.row_y[r1]
        for key, col in m.columns.items():
            x, w = m.col_x[col], m.rect(r1, col)[2]
            value = row.get(key)
            if not value:
                continue
            if key.startswith("sig"):
                reader = _signature_reader(value)
                if reader:
                    _draw_image(c, reader, x, y, w, h)
            else:
                _draw_text(c, value, x, y, w, h, "Helvetica", 9.5, black)
    c.restoreState()


def build_pdf(header, rows, footer=""):
    """header: {'operator', 'name'}; rows: dict con date_out, key, sig_out, sig_in, date_in."""
    m = model()
    per_page = max(len(m.body), 1)
    pages = [rows[i:i + per_page] for i in range(0, len(rows), per_page)] or [[]]
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=m.pagesize)
    c.setTitle(footer or "Ricevuta chiavi")
    c.setAuthor(header.get("operator") or "")
    stamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    for n, chunk in enumerate(pages, 1):
        _draw_page(c, m, header, chunk)
        c.setFont("Helvetica", 7.5)
        c.setFillColor(Color(.45, .45, .45))
        text = f"© {datetime.now().year} {config.COPYRIGHT_OWNER} · {footer} · generato il {stamp}" + (f" · pagina {n} di {len(pages)}" if len(pages) > 1 else "")
        c.drawRightString(m.pagesize[0] - 18, m.margin_bottom / 2, text.strip(" ·"))
        c.showPage()
    c.save()
    return buf.getvalue()
